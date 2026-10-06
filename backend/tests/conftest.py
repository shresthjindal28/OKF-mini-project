import io
from pathlib import Path

import pytest
from docx import Document
from reportlab.pdfgen.canvas import Canvas

from app.core.config import Settings


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        database_url="postgresql+psycopg://test:test@localhost/okf_test",
        hf_token="test-token",
        storage_path=tmp_path,
        hf_max_retries=0,
    )


@pytest.fixture
def pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.setTitle("Research report")
    canvas.setAuthor("Test Author")
    canvas.drawString(72, 750, "Knowledge portability and semantic indexing.")
    canvas.save()
    return buffer.getvalue()


@pytest.fixture
def docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.core_properties.title = "Research report"
    document.core_properties.author = "Test Author"
    document.add_heading("Knowledge portability", level=1)
    document.add_paragraph("Semantic indexing preserves portable knowledge.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Format"
    table.cell(1, 1).text = "OKF"
    document.save(buffer)
    return buffer.getvalue()
